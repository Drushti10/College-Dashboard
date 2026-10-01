"""
Script to create sample PDF and DOCX files for testing DocuMind AI.
"""

import os
from pathlib import Path

SAMPLE_DIR = Path(__file__).resolve().parent / "data" / "sample_document"
SAMPLE_DIR.mkdir(parents=True, exist_ok=True)


def create_minimal_pdf(filepath: Path):
    """
    Creates a valid 4-page PDF document about Artificial Intelligence & Machine Learning
    using standard PDF 1.4 objects without external dependencies.
    """
    pages_content = [
        # Page 1
        "ARTIFICIAL INTELLIGENCE OVERVIEW\n\n"
        "Artificial Intelligence (AI) refers to the simulation of human intelligence in machines.\n"
        "The goal of AI is to develop systems capable of performing tasks that typically require human cognition,\n"
        "such as visual perception, speech recognition, decision-making, and language translation.\n\n"
        "Machine Learning is a subset of AI that focuses on building applications that learn from data and improve\n"
        "their accuracy over time without being explicitly programmed to do so.",

        # Page 2
        "SUPERVISED AND UNSUPERVISED LEARNING\n\n"
        "Supervised learning algorithms are trained using labeled datasets, where the desired output is known.\n"
        "Common algorithms include Linear Regression, Support Vector Machines, and Decision Trees.\n"
        "Supervised learning is widely used for classification and regression tasks like spam detection and price prediction.\n\n"
        "Unsupervised learning uses unlabeled data to discover hidden patterns and clusters without human guidance.\n"
        "Examples include K-Means Clustering and Principal Component Analysis (PCA).",

        # Page 3
        "DEEP LEARNING AND NEURAL NETWORKS\n\n"
        "Deep Learning is a specialized subfield of Machine Learning based on Artificial Neural Networks with representation learning.\n"
        "The term deep refers to the number of layers in the network.\n"
        "A typical deep neural network consists of an input layer, multiple hidden layers, and an output layer.\n\n"
        "Convolutional Neural Networks (CNNs) excel at image processing, while Recurrent Neural Networks (RNNs) and\n"
        "Transformers are designed for sequential data and natural language processing.",

        # Page 4
        "RETRIEVAL-AUGMENTED GENERATION (RAG)\n\n"
        "Retrieval-Augmented Generation (RAG) is an architectural framework that improves the accuracy of Large Language Models (LLMs).\n"
        "Instead of relying solely on training weights, RAG retrieves relevant facts from an external knowledge base\n"
        "or private document collection before generating a response.\n\n"
        "Key advantages of RAG include:\n"
        "1. Reduction in hallucination by grounding responses in verified source texts.\n"
        "2. Access to domain-specific and up-to-date data without expensive fine-tuning.\n"
        "3. Verifiable source citations and page transparency for users."
    ]

    objects = []
    # 1: Catalog, 2: Outlines, 3: Pages
    catalog_id = 1
    outlines_id = 2
    pages_id = 3
    font_id = 4
    
    current_obj_id = 5
    page_ids = []
    content_ids = []

    for _ in pages_content:
        page_ids.append(current_obj_id)
        current_obj_id += 1
        content_ids.append(current_obj_id)
        current_obj_id += 1

    # Font object
    font_obj = f"{font_id} 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
    
    # Page objects and Content streams
    page_and_stream_objs = []
    for i, text in enumerate(pages_content):
        pid = page_ids[i]
        cid = content_ids[i]
        
        # Format text lines for PDF stream
        lines = text.split("\n")
        stream_cmds = ["BT", "/F1 12 Tf", "50 720 Td", "16 TL"]
        for line in lines:
            safe_line = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            if safe_line:
                stream_cmds.append(f"({safe_line}) '")
            else:
                stream_cmds.append("T*")
        stream_cmds.append("ET")
        stream_data = "\n".join(stream_cmds)
        
        p_obj = (
            f"{pid} 0 obj\n"
            f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 612 792] "
            f"/Contents {cid} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>\n"
            f"endobj\n"
        )
        c_obj = (
            f"{cid} 0 obj\n"
            f"<< /Length {len(stream_data)} >>\n"
            f"stream\n{stream_data}\nendstream\n"
            f"endobj\n"
        )
        page_and_stream_objs.extend([p_obj, c_obj])

    kids_str = " ".join([f"{p} 0 R" for p in page_ids])
    pages_obj = (
        f"{pages_id} 0 obj\n"
        f"<< /Type /Pages /Kids [{kids_str}] /Count {len(page_ids)} >>\n"
        f"endobj\n"
    )
    catalog_obj = (
        f"{catalog_id} 0 obj\n"
        f"<< /Type /Catalog /Pages {pages_id} 0 R >>\n"
        f"endobj\n"
    )
    outlines_obj = (
        f"{outlines_id} 0 obj\n"
        f"<< /Type /Outlines /Count 0 >>\n"
        f"endobj\n"
    )

    all_objs = [catalog_obj, outlines_obj, pages_obj, font_obj] + page_and_stream_objs
    
    # Build complete PDF binary
    header = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n"
    body = b""
    xref_offsets = [0]
    
    for obj_str in all_objs:
        xref_offsets.append(len(header) + len(body))
        body += obj_str.encode("latin1")
        
    xref_pos = len(header) + len(body)
    total_objs = len(xref_offsets)
    
    xref = f"xref\n0 {total_objs}\n0000000000 65535 f \n"
    for off in xref_offsets[1:]:
        xref += f"{off:010d} 00000 n \n"
        
    trailer = (
        f"trailer\n"
        f"<< /Size {total_objs} /Root {catalog_id} 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    )
    
    pdf_bytes = header + body + xref.encode("latin1") + trailer.encode("latin1")
    with open(filepath, "wb") as f:
        f.write(pdf_bytes)
    print(f"Created sample PDF: {filepath}")


def create_sample_docx(filepath: Path):
    """Creates a sample DOCX document."""
    try:
        import docx
        doc = docx.Document()
        doc.add_heading("Artificial Intelligence and Machine Learning Handbook", level=0)
        
        doc.add_heading("1. Machine Learning Fundamentals", level=1)
        doc.add_paragraph(
            "Machine Learning enables computers to learn from data patterns rather than explicitly coded logic. "
            "It is divided into supervised learning, unsupervised learning, and reinforcement learning."
        )
        doc.add_paragraph(
            "Supervised learning requires labeled training examples, whereas unsupervised learning identifies intrinsic data clustering."
        )
        
        doc.add_heading("2. Neural Networks & Transformer Architectures", level=1)
        doc.add_paragraph(
            "Artificial Neural Networks are modeled on human neurobiology with interconnected nodes organized in sequential layers."
        )
        doc.add_paragraph(
            "The Transformer architecture, introduced by Vaswani et al. in 2017, relies on self-attention mechanisms "
            "to process sequential data concurrently, forming the foundation of modern Large Language Models."
        )
        
        doc.add_heading("3. RAG Architecture and Vector Search", level=1)
        doc.add_paragraph(
            "Retrieval-Augmented Generation (RAG) pairs semantic document retrieval with generative language models. "
            "By converting documents into high-dimensional vector embeddings stored in vector databases like FAISS, "
            "systems can accurately retrieve semantically relevant context chunks and cite their exact document origins."
        )
        
        doc.save(filepath)
        print(f"Created sample DOCX: {filepath}")
    except ImportError:
        print("python-docx not yet installed. Will create DOCX after pip install completes.")


if __name__ == "__main__":
    pdf_path = SAMPLE_DIR / "sample_ai_overview.pdf"
    docx_path = SAMPLE_DIR / "sample_ai_overview.docx"
    create_minimal_pdf(pdf_path)
    create_sample_docx(docx_path)
